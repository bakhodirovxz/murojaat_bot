"""Arizalardan «Qabul jadvali» ko'rinishidagi Excel fayl yaratish.

Ustun kengligi va qator balandligi tarkibga qarab avtomatik hisoblanadi,
shuning uchun uzun manzil yoki murojaat matni kesilib qolmaydi.
"""

import math

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import db

TITLE = "Qabul jadvali"

HEADERS = (
    "№",
    "Fuqaro F.I.Sh.",
    "Pasport seriyasi",
    "Fuqaro JSHSHIR raqami",
    "Tug'ilgan yili va sanasi",
    "Yashash manzili (to'liq)",
    "Telefon raqami (qo'shimcha)",
    "Murojaat mazmuni",
    "Qabul vaqti",
    "Holati",
    "Berilgan javob",
)

# Ustun cheksiz kengayib ketmasligi uchun yuqori chegara (belgi hisobida).
MAX_WIDTHS = (7, 34, 18, 22, 18, 46, 28, 62, 18, 20, 62)

LINE_HEIGHT = 15.0
MAX_ROW_HEIGHT = 400.0

# Excel kenglikni piksel bilan o'lchaydi, biz esa belgi soni bilan: keng harflar
# (М, Ш, w) uchun zaxira qoldiramiz, aks holda oxirgi qator kesilib qolishi mumkin.
WIDTH_SAFETY = 0.9

_THIN = Side(style="thin", color="999999")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def format_phones(phone: str, phone_extra: str) -> str:
    """Asosiy va qo'shimcha raqamni bitta katakka joylaydi."""
    if phone_extra:
        return f"{phone} ({phone_extra})"
    return phone or ""


def pretty_datetime(value: str) -> str:
    """`2026-09-04T15:23:01` → `04.09.2026 15:23`."""
    if not value or "T" not in value:
        return value or ""
    day_part, time_part = value.split("T", 1)
    try:
        year, month, day = day_part.split("-")
    except ValueError:
        return value
    return f"{day}.{month}.{year} {time_part[:5]}"


def row_values(row: dict) -> tuple:
    """Baza yozuvidan jadval qatorini yasaydi."""
    return (
        row["id"],
        row["fio"],
        row["passport"],
        row["jshshir"],
        row["birth_date"],
        row["address"],
        format_phones(row["phone"], row.get("phone_extra") or ""),
        row["message"],
        pretty_datetime(row.get("created_at", "")),
        db.STATUS_LABELS.get(row.get("status") or db.STATUS_NEW, "—"),
        row.get("last_answer") or "",
    )


def _longest_word(text: str) -> int:
    return max((len(word) for word in text.split()), default=0)


def compute_widths(rows) -> list:
    """Har bir ustunga sarlavha va tarkibga mos kenglik tanlaydi."""
    widths = []
    for index, header in enumerate(HEADERS):
        # Sarlavha ikki qatorga o'ralishi mumkin, lekin eng uzun so'z sig'sin.
        header_need = max(5, math.ceil(len(header) / 2) + 2, _longest_word(header) + 2)
        content_need = header_need
        for row in rows:
            text = str(row[index] if row[index] is not None else "")
            longest_line = max((len(line) for line in text.split("\n")), default=0)
            content_need = max(content_need, longest_line + 2)
        widths.append(float(min(MAX_WIDTHS[index], content_need)))
    return widths


def _wrapped_line_count(text: str, width: float) -> int:
    """Excel matnni so'z bo'yicha o'raydi — qator sonini shunga mos sanaymiz."""
    usable = max(int(width * WIDTH_SAFETY) - 1, 4)
    lines = 0
    for segment in str(text).split("\n"):
        words = segment.split()
        if not words:
            lines += 1
            continue
        count = 1
        filled = 0
        for word in words:
            needed = len(word) if filled == 0 else filled + 1 + len(word)
            if needed <= usable:
                filled = needed
                continue
            count += 1
            filled = len(word)
            while filled > usable:  # bitta so'z ustunga sig'masa bo'linadi
                count += 1
                filled -= usable
        lines += count
    return max(1, lines)


def compute_row_height(values, widths) -> float:
    """Eng ko'p o'ralgan katakka qarab qator balandligini beradi."""
    lines = 1
    for index, value in enumerate(values):
        text = "" if value is None else str(value)
        lines = max(lines, _wrapped_line_count(text, widths[index]))
    return min(MAX_ROW_HEIGHT, lines * LINE_HEIGHT + 6)


def build_workbook(rows) -> Workbook:
    table = [row_values(row) for row in rows]
    widths = compute_widths(table)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Qabul"

    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEADERS))
    title_cell = sheet.cell(row=1, column=1, value=TITLE)
    title_cell.font = Font(bold=True, size=14)
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 26

    header_fill = PatternFill("solid", fgColor="DDEBF7")
    for index, header in enumerate(HEADERS, start=1):
        cell = sheet.cell(row=2, column=index, value=header)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _BORDER
        sheet.column_dimensions[get_column_letter(index)].width = widths[index - 1]
    sheet.row_dimensions[2].height = compute_row_height(HEADERS, widths)

    for offset, values in enumerate(table):
        excel_row = 3 + offset
        for index, value in enumerate(values, start=1):
            cell = sheet.cell(row=excel_row, column=index, value=value)
            cell.alignment = Alignment(
                vertical="top",
                horizontal="center" if index in (1, 3, 4, 5) else "left",
                wrap_text=True,
            )
            cell.border = _BORDER
        sheet.row_dimensions[excel_row].height = compute_row_height(values, widths)

    sheet.freeze_panes = "A3"
    sheet.auto_filter.ref = f"A2:{get_column_letter(len(HEADERS))}{2 + len(table)}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    return workbook


def save_workbook(rows, path: str) -> str:
    build_workbook(rows).save(path)
    return path
