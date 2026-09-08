import pytest

import db
import export


SAMPLE = {
    "fio": "Karimov Aziz Baxtiyorovich",
    "passport": "AA1234567",
    "jshshir": "31234567890123",
    "birth_date": "15.03.1990",
    "address": "Toshkent sh., Chilonzor t., 5-uy",
    "phone": "+998901234567",
    "phone_extra": "+998971234567",
    "message": "Uy-joy hujjatlari masalasida murojaat.",
    "tg_user_id": 111,
    "tg_username": "aziz",
}


@pytest.fixture()
def database(tmp_path):
    path = str(tmp_path / "test.db")
    db.init_db(path)
    return path


class TestDatabase:
    def test_init_is_idempotent(self, database):
        db.init_db(database)
        assert db.count_applications(database) == 0

    def test_add_returns_sequential_ids(self, database):
        assert db.add_application(database, SAMPLE) == 1
        assert db.add_application(database, SAMPLE) == 2
        assert db.count_applications(database) == 2

    def test_fetch_returns_stored_values(self, database):
        db.add_application(database, SAMPLE)
        rows = db.fetch_applications(database)
        assert len(rows) == 1
        row = rows[0]
        for key, value in SAMPLE.items():
            assert row[key] == value
        assert row["created_at"]

    def test_date_filter(self, database):
        db.add_application(database, SAMPLE)
        assert db.fetch_applications(database, "1990-01-01", "1990-12-31") == []
        assert len(db.fetch_applications(database, "1990-01-01", "2099-12-31")) == 1

    def test_last_created_at(self, database):
        assert db.last_created_at(database) is None
        db.add_application(database, SAMPLE)
        assert db.last_created_at(database) is not None


class TestExport:
    def test_format_phones(self):
        assert export.format_phones("+998901234567", "+998971234567") == (
            "+998901234567 (+998971234567)"
        )
        assert export.format_phones("+998901234567", "") == "+998901234567"

    def test_pretty_datetime(self):
        assert export.pretty_datetime("2026-09-04T15:23:01") == "04.09.2026 15:23"
        assert export.pretty_datetime("") == ""

    def test_workbook_layout(self, database, tmp_path):
        db.add_application(database, SAMPLE)
        rows = db.fetch_applications(database)
        path = str(tmp_path / "out.xlsx")
        export.save_workbook(rows, path)

        from openpyxl import load_workbook

        sheet = load_workbook(path).active
        assert sheet.cell(row=1, column=1).value == export.TITLE
        assert [c.value for c in sheet[2]] == list(export.HEADERS)
        assert sheet.cell(row=3, column=1).value == 1
        assert sheet.cell(row=3, column=2).value == SAMPLE["fio"]
        assert sheet.cell(row=3, column=7).value == "+998901234567 (+998971234567)"


class TestAutoSizing:
    def _row(self, **overrides):
        row = dict(SAMPLE, id=1, created_at="2026-09-04T17:00:00")
        row.update(overrides)
        return row

    def test_width_fits_content_but_stays_within_limit(self):
        short = export.compute_widths([export.row_values(self._row())])
        long_message = export.compute_widths(
            [export.row_values(self._row(message="so'z " * 300))]
        )
        message_column = export.HEADERS.index("Murojaat mazmuni")
        assert short[message_column] < long_message[message_column]
        for index, width in enumerate(long_message):
            assert width <= export.MAX_WIDTHS[index]

    def test_width_never_hides_content_that_fits(self):
        address = "Toshkent sh., Chilonzor tumani, 5-uy"
        widths = export.compute_widths([export.row_values(self._row(address=address))])
        address_column = export.HEADERS.index("Yashash manzili (to'liq)")
        assert widths[address_column] >= len(address) + 2

    def test_line_count_wraps_on_word_boundaries(self):
        # 10 belgilik ustunda "aaaa bbbb cccc" ikki qatorga tushadi.
        assert export._wrapped_line_count("aaaa bbbb cccc", 12) == 2
        assert export._wrapped_line_count("qisqa", 10) == 1
        assert export._wrapped_line_count("a" * 40, 10) >= 4

    def test_row_height_grows_with_text_and_is_capped(self):
        widths = export.compute_widths([export.row_values(self._row())])
        short = export.compute_row_height(export.row_values(self._row()), widths)
        tall = export.compute_row_height(
            export.row_values(self._row(message="so'z " * 400)), widths
        )
        assert tall > short
        assert tall <= export.MAX_ROW_HEIGHT

    def test_workbook_sets_width_and_height_for_every_row(self, database, tmp_path):
        db.add_application(database, SAMPLE)
        db.add_application(
            database, dict(SAMPLE, message="Uzun murojaat matni. " * 40, address="A" * 200)
        )
        rows = db.fetch_applications(database)
        path = str(tmp_path / "sized.xlsx")
        export.save_workbook(rows, path)

        from openpyxl import load_workbook
        from openpyxl.utils import get_column_letter

        sheet = load_workbook(path).active
        for index in range(1, len(export.HEADERS) + 1):
            width = sheet.column_dimensions[get_column_letter(index)].width
            assert 0 < width <= export.MAX_WIDTHS[index - 1]
        for row_number in range(2, sheet.max_row + 1):
            assert sheet.row_dimensions[row_number].height > 0
            assert sheet.cell(row=row_number, column=8).alignment.wrap_text is True
        # Uzun matnli qator qisqasidan baland bo'lishi kerak.
        assert sheet.row_dimensions[4].height > sheet.row_dimensions[3].height
